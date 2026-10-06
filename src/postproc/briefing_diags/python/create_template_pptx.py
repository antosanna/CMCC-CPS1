###################################################################################################
## Script:  create_template_pptx.py                                                              ##
## Author:  Brandon J Fisel                                                                      ##
## Summary: Replaces template pptx images with image placeholders.                               ##
## Description: Reads in pptx template files: orig_template_summer.pptx and                      ##
##          orig_template_winter.pptx saved in Powerpoint or Keynote. The original images are    ##
##          replaced by image placeholders, and those placeholders are then replaced with the    ##
##          new images by modify_template_dev.py. This script writes the new pptx files:         ##
##          template_summer.pptx and template_winter.pptx within the pptx templates directory.   ##
##          Use month=6 (summer) or month=1 (winter) to create the seasonal pptx template files. ##
##          Run as: python3 create_template_pptx.py --input pptx/orig_template_summer.pptx \     ##
##                  --output pptx/template_summer.pptx --month 6                                 ##
##          Validate as: python3 create_template_pptx.py --input pptx/template_summer.pptx \     ##
##                       --validate                                                              ##
## Creation Date: 16/09/2026                                                                     ##
## Revision Date: 21/09/2026                                                                     ##
## Revised By: BJF                                                                               ##
###################################################################################################

import argparse
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN

try:
    from modify_template_dev import get_named_slide_image_mapping
except ImportError:
    print("Error: Couldn't import 'get_named_slide_image_mapping' from modify_template_dev.py.")
    print("Ensure modify_template_dev.py is in the current working directory.")
    exit(1)


def get_shape_alt_text(shape):
    # Retrieves alt text (description) across different python-pptx shape wrappers
    try:
        if hasattr(shape, 'alt_text') and shape.alt_text:
            return shape.alt_text.strip()
    except Exception:
        pass

    try:
        element = shape._element
        for child in element.iter():
            if child.tag.endswith('cNvPr'):
                return child.attrib.get('descr', '').strip()
    except Exception:
        pass

    return ""


def set_shape_alt_text(shape, text):
    # Sets alt text in the underlying cNvPr XML element"
    try:
        shape.alt_text = text
    except Exception:
        pass

    try:
        element = shape._element
        for child in element.iter():
            if child.tag.endswith('cNvPr'):
                child.attrib['descr'] = text
                break
    except Exception:
        pass


def validate_presentation_tags(pptx_path):
    # Scans a briefing pptx template file and prints out shape names and alt text per slide
    print(f"\n==========================================")
    print(f" VALIDATING SHAPE TAGS: {pptx_path}")
    print(f"==========================================\n")

    prs = Presentation(pptx_path)

    for slide_idx, slide in enumerate(prs.slides):
        print(f"--- Slide {slide_idx} ---")
        tagged_shapes_found = 0

        for shape_idx, shape in enumerate(slide.shapes):
            alt_text = get_shape_alt_text(shape)
            shape_name = shape.name.strip() if shape.name else ""

            sample_text = ""
            if shape.has_text_frame and shape.text_frame.text:
                clean_text = shape.text_frame.text.replace("\n", " ").strip()
                sample_text = f" | Text: '{clean_text[:30]}...'" if len(clean_text) > 30 else f" | Text: '{clean_text}'"

            if alt_text or "IMG_" in shape_name:
                tagged_shapes_found += 1
                print(f"  [Layer {shape_idx}] Tag: '{alt_text or shape_name}' (Alt Text: '{alt_text}', Name: '{shape_name}'){sample_text}")
            else:
                print(f"  [Layer {shape_idx}] Shape: {shape.shape_type}{sample_text}")

        if tagged_shapes_found == 0:
            print("  (No image placeholders found on this slide)")
        print()


def convert_images_to_placeholders(input_pptx, output_pptx, month=9):
    # Converts image shapes into borderless placeholder shapes
    prs = Presentation(input_pptx)
    image_mappings = get_named_slide_image_mapping(month)

    total_placeholders_created = 0

    for slide_idx, shape_map in image_mappings.items():
        if slide_idx >= len(prs.slides):
            continue

        slide = prs.slides[slide_idx]

        pictures_to_replace = []
        for shape in slide.shapes:
            if shape.shape_type == 13:
                pictures_to_replace.append(shape)

        identifiers = list(shape_map.keys())

        for idx, pic_shape in enumerate(pictures_to_replace):
            shape_id = get_shape_alt_text(pic_shape) or pic_shape.name

            if shape_id not in shape_map and idx < len(identifiers):
                shape_id = identifiers[idx]

            if shape_id not in shape_map:
                print(f"Slide {slide_idx}: Skipping unmapped image '{pic_shape.name}'.")
                continue

            left, top = pic_shape.left, pic_shape.top
            width, height = pic_shape.width, pic_shape.height

            # Remove old picture element
            pic_element = pic_shape._element
            sp_tree = pic_element.getparent()
            shape_index = sp_tree.index(pic_element)
            sp_tree.remove(pic_element)

            # Add placeholder shape
            placeholder = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, left, top, width, height
            )

            # Placeholder background and border
            placeholder.fill.solid()
            placeholder.fill.fore_color.rgb = RGBColor(235, 235, 240)
            placeholder.line.fill.background()

            # Set placeholder label text
            text_frame = placeholder.text_frame
            text_frame.word_wrap = True
            p = text_frame.paragraphs[0]
            p.text = shape_id
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = RGBColor(90, 90, 100)

            # Set alt text in both Python wrapper and XML tag
            set_shape_alt_text(placeholder, shape_id)

            # Keep placeholder at original layer position
            ph_element = placeholder._element
            sp_tree.remove(ph_element)
            sp_tree.insert(shape_index, ph_element)

            total_placeholders_created += 1
            print(f"Slide {slide_idx}: Replaced image with placeholder '{shape_id}'")

    prs.save(output_pptx)
    print(f"\nDone! Created {total_placeholders_created} placeholder.")
    print(f"Saved template to: {output_pptx}")


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Convert pptx images to placeholder shapes or validate pptx shapes.")
    parser.add_argument('--input', '-i', required=True, help="Path to input original presentation.")
    parser.add_argument('--output', '-o', help="Path to save template with shape placeholders.")
    parser.add_argument('--month', '-m', type=int, default=9, help="Month index for seasonal mappings (1-12).")
    parser.add_argument('--validate', '-v', action='store_true', help="Run validation to inspect shape placeholders.")

    args = parser.parse_args()

    if args.validate:
        validate_presentation_tags(args.input)
    else:
        if not args.output:
            print("Error: --output / -o parameter is required when creating placeholders!")
            exit(1)
        convert_images_to_placeholders(args.input, args.output, month=args.month)
