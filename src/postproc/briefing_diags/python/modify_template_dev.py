###################################################################################################
## Script:  modify_template_dev.py                                                               ##
## Author:  Brandon J Fisel                                                                      ##
## Summary: Creates the briefing presentation from required images.                              ##
## Description: Replaces image placeholders with required briefing images based on their alt     ##
##          text name. Slide text is also replaced where applicable. The briefing presentation   ##
##          pptx is then created in $SCRATCHDIR. Some of the functions in this script are used   ##
##          by create_template_pptx, so revisions may need to be done to both scripts.           ##
## Creation Date: 16/09/2026                                                                     ##
## Revision Date: 25/09/2026                                                                     ##
## Revised By: BJF                                                                               ##
###################################################################################################

import argparse
import sys
from datetime import datetime
from pathlib import Path
from PIL import Image as PILImage
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def get_shape_alt_text(shape):
    # Retrieve alt text across different python-pptx shape wrappers
    try:
        if hasattr(shape, 'alt_text') and shape.alt_text:
            return shape.alt_text.strip()
    except Exception:
        pass

    try:
        element = shape._element
        for child in element.iter():
            if child.tag.endswith('cNvPr'):
                val = child.attrib.get('descr', '').strip()
                if val:
                    return val
    except Exception:
        pass

    return ""


def get_shape_identifier(shape):
    # Extracts tag identifier from alt text, shape name or paragraph text
    alt_text = get_shape_alt_text(shape)
    if alt_text:
        return alt_text

    if shape.name and "IMG_" in shape.name:
        return shape.name.strip()

    if shape.has_text_frame and shape.text_frame.text:
        text_content = shape.text_frame.text.strip()
        if text_content.startswith("IMG_"):
            return text_content

    return ""


def fit_image_in_bounds(image_path, target_left, target_top, target_width, target_height):
    # Calculate dimensions and offsets to center-fit images inside the bounding box
    with PILImage.open(image_path) as img:
        img_width, img_height = img.size

    img_aspect = img_width / img_height
    target_aspect = target_width / target_height

    if img_aspect > target_aspect:
        # Fitting to width
        new_width = target_width
        new_height = int(target_width / img_aspect)
        new_left = target_left
        new_top = target_top + (target_height - new_height) // 2
    else:
        # Fitting to height
        new_height = target_height
        new_width = int(target_height * img_aspect)
        new_top = target_top
        new_left = target_left + (target_width - new_width) // 2

    return new_left, new_top, new_width, new_height


def replace_named_images_on_slide(slide, image_map, img_desc_vars):
    # Replace shapes on slides matching the shape ID with incoming images, maintaining
    # aspect-fit boundaries and z-order layer positions
    targets = []

    # Collect all matching shapes on the slide
    for shape in slide.shapes:
        shape_identifier = get_shape_identifier(shape)

        if shape_identifier in image_map:
            path_template = image_map[shape_identifier]
            # Format pattern with description identifiers
            formatted_path_str = path_template.format(**img_desc_vars)
            img_path = Path(formatted_path_str).resolve()

            if not img_path.is_file():
                print(f"Warning: Image file not found: {img_path}. Skipping shape '{shape_identifier}'.")
                continue

            targets.append((shape, shape_identifier, img_path))

    # Swap out shapes for incoming images
    for shape, shape_identifier, img_path in targets:
        left, top = shape.left, shape.top
        width, height = shape.width, shape.height

        new_left, new_top, new_w, new_h = fit_image_in_bounds(
            img_path, left, top, width, height
        )

        # Normal picture shape - swap directly
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            with open(img_path, 'rb') as f:
                shape.image.blob = f.read()
            print(f"Updated existing picture for '{shape_identifier}'")

        # Swap placeholder with incoming image at XML index
        else:
            old_element = shape._element
            sp_tree = old_element.getparent()
            target_index = sp_tree.index(old_element)

            # Insert new picture shape
            new_pic = slide.shapes.add_picture(str(img_path), new_left, new_top, width=new_w, height=new_h)
            pic_element = new_pic._element

            # Shift new picture into original and delete placeholder
            sp_tree.remove(pic_element)
            sp_tree.insert(target_index, pic_element)
            sp_tree.remove(old_element)

            print(f"Replaced placeholder '{shape_identifier}' with image: {img_path.name}")


def replace_text_in_presentation(prs, replacements):
    # Replace text placeholders across all slides and retain font formatting
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                full_text = "".join(run.text for run in paragraph.runs)

                # Replace text if placeholder string exists
                for key, val in replacements.items():
                    if key in full_text:
                        full_text = full_text.replace(key, str(val))

                # Reassign text to first run and clear subsequent runs to keep the original style
                if paragraph.runs:
                    paragraph.runs[0].text = full_text
                    for run in paragraph.runs[1:]:
                        run.text = ""


def get_named_slide_image_mapping(month):
    # Returns the slide-index mapped to named shape bindings
    # Format: slide_index: { 'SHAPE_NAME': 'image_file.png' }
    mapping = {
        2: {'IMG_SIGNIF_EVENTS': '{IMG_DNLD_DIR}/signif_events.png'},
        3: {'IMG_NINO_TS': '{IMG_DNLD_DIR}/nino_timeseries.png'},
        4: {
            'IMG_SST_BOTTOM': '{IMG_DNLD_DIR}/sst_anomalies_bottom.png',
            'IMG_SST_TOP': '{IMG_DNLD_DIR}/sst_anomalies_top.png'
        },
        5: {
            'IMG_ARCTIC_MAP': '{IMG_DNLD_DIR}/map_arctic.png',
            'IMG_ANTARCTIC_MAP': '{IMG_DNLD_DIR}/map_antarctic.png',
            'IMG_ANTARCTIC_TS': '{IMG_DNLD_DIR}/timesrs_antarctic.png',
            'IMG_ARCTIC_TS': '{IMG_DNLD_DIR}/timesrs_arctic.png'
        },
        6: {
            'IMG_T2M_EU': '{IMG_DNLD_DIR}/t2m_c3s_climbull_eu.png',
            'IMG_T2M_GLO': '{IMG_DNLD_DIR}/t2m_c3s_climbull_glo.png'
        },
        7: {
            'IMG_HYDRO_C3S': '{IMG_DNLD_DIR}/hydro_c3s_climbull.png',
            'IMG_MSLP_ANOM': '{IMG_CIRC_DIR}/mslp_anom_{yyyym1}{mm1}.png',
            'IMG_Z500_ANOM': '{IMG_CIRC_DIR}/z500_anom_{yyyym1}{mm1}.png',
            'IMG_U250_ANOM': '{IMG_CIRC_DIR}/u250_anom_{yyyym1}{mm1}.png'
        },
        9: {
            'IMG_COLORBAR': '{IMG_EVAL_DIR}/eval_colorbar.png',
            'IMG_PRECIP_EVAL_EU': '{IMG_EVAL_DIR}/precip_ano_verification_europe_{yyyym2}{mm2}_l1_monthly_nobar.png',
            'IMG_T2M_EVAL_EU': '{IMG_EVAL_DIR}/t2m_ano_verification_europe_{yyyym2}{mm2}_l1_monthly_nobar.png'
        },
        10: {
            'IMG_COLORBAR': '{IMG_EVAL_DIR}/eval_colorbar.png',
            'IMG_T2M_EVAL_GLOB': '{IMG_EVAL_DIR}/t2m_ano_verification_global_{yyyym2}{mm2}_l1_monthly_nobar.png',
            'IMG_PRECIP_EVAL_GLOB': '{IMG_EVAL_DIR}/precip_ano_verification_global_{yyyym2}{mm2}_l1_monthly_nobar.png'
        },
        12: {'IMG_NINO_EVAL': '{IMG_EVAL_DIR}/Nino3.4_verification_{yyyym1}{mm1}.png'},
        13: {'IMG_SIE_EVAL': '{IMG_SIE_DIR}/NH_SIE_{sm1forey}{sm1forem}.png'},
        14: {
            'IMG_COLORBAR': '{IMG_EVAL_DIR}/eval_colorbar.png',
            'IMG_PRECIP_EVAL_GLOB_SEASON': '{IMG_EVAL_DIR}/precip_ano_verification_global_{sm1forey}{sm1forem}_l1_seasonal_nobar.png',
            'IMG_T2M_EVAL_GLOB_SEASON': '{IMG_EVAL_DIR}/t2m_ano_verification_global_{sm1forey}{sm1forem}_l1_seasonal_nobar.png'
        },
        15: {
            'IMG_COLORBAR': '{IMG_EVAL_DIR}/eval_colorbar.png',
            'IMG_PRECIP_EVAL_EU_SEASON': '{IMG_EVAL_DIR}/precip_ano_verification_europe_{sm1forey}{sm1forem}_l1_seasonal_nobar.png',
            'IMG_T2M_EVAL_EU_SEASON': '{IMG_EVAL_DIR}/t2m_ano_verification_europe_{sm1forey}{sm1forem}_l1_seasonal_nobar.png'
        },
        17: {'IMG_C3S_MME_NINO': '{IMG_DNLD_DIR}/c3s_mme_nino.png'},
        18: {
            'IMG_SST_NINO_PROB': '{IMG_DIAG_DIR}/sst_Nino3.4_strength_prob_{yyyy}_{mm}.png',
            'IMG_SST_NINO_MEM': '{IMG_DIAG_DIR}/sst_Nino3.4_mem_{yyyy}_{mm}.png'
        },
        19: {'IMG_TPROF_PAC_TROP': '{IMG_WEB_DIR}/temperature_pac_trop_ensmean_{yyyy}_{mm}.gif'},
        23: {
            'IMG_SSTO_FORE_CMCC': '{IMG_DNLD_DIR}/cmcc_sst_fore_global.png',
            'IMG_SSTO_FORE_MM': '{IMG_DNLD_DIR}/c3s_mm_ssto_fore_glob.png'
        },
        24: {
            'IMG_MSLP_FORE_CMCC': '{IMG_DNLD_DIR}/cmcc_mslp_fore_global.png',
            'IMG_MSLP_FORE_MM': '{IMG_DNLD_DIR}/c3s_mm_mslp_fore_glob.png'
        },
        25: {
            'IMG_Z500_FORE_CMCC': '{IMG_DNLD_DIR}/cmcc_hgt500_fore_global.png',
            'IMG_Z500_FORE_MM': '{IMG_DNLD_DIR}/c3s_mm_z500_fore_glob.png'
        },
        26: {
            'IMG_RAIN_FORE_CMCC': '{IMG_DNLD_DIR}/cmcc_precip_fore_global.png',
            'IMG_RAIN_FORE_MM': '{IMG_DNLD_DIR}/c3s_mm_rain_fore_glob.png'
        },
        27: {
            'IMG_2MTM_FORE_CMCC': '{IMG_DNLD_DIR}/cmcc_t2m_fore_global.png',
            'IMG_2MTM_FORE_MM': '{IMG_DNLD_DIR}/c3s_mm_2mtm_fore_glob.png'
        },
        28: {
            'IMG_Z500_EURO_CMCC': '{IMG_DNLD_DIR}/cmcc_hgt500_fore_Europe.png',
            'IMG_MSLP_EURO_CMCC': '{IMG_DNLD_DIR}/cmcc_mslp_fore_Europe.png'
        },
        29: {
            'IMG_2MTM_EURO_CMCC': '{IMG_DNLD_DIR}/cmcc_t2m_fore_Europe.png',
            'IMG_RAIN_EURO_CMCC': '{IMG_DNLD_DIR}/cmcc_precip_fore_Europe.png'
        },
        30: {'IMG_FORECAST_SUMMARY': '{IMG_DNLD_DIR}/Europe_summary_{yyyy}_{mm}_l1.png'}
    }

    # Seasonal variation rules
    is_winter = (month < 4 or month >= 10)
    if is_winter:
        mapping[20] = {'IMG_U10HPA_CMCC': '{IMG_DNLD_DIR}/U10hPa_cmcc_fore_probs_lt0.png'}
        mapping[21] = {'IMG_U10HPA_PROBS': '{IMG_DNLD_DIR}/U10hPa_fore_probs_lt0.png'}
    else:
        mapping[21] = {
            'IMG_IOD_MEM': '{IMG_WEB_DIR}/sst_IOD_mem_{yyyy}_{mm}.png',
            'IMG_IOD_PROB': '{IMG_WEB_DIR}/sst_IOD_prob_{yyyy}_{mm}.png'
        }
        mapping[32] = {
            'IMG_FORECAST_WAMI': '{IMG_WAMI_DIR}/WAMI_forecast_{yyyy}{mm}.png',
            'IMG_FORECAST_WAMI_PROB': '{IMG_WAMI_DIR}/WAMI_probability_{yyyy}{mm}.png'
        }

    return mapping

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Populate powerpoint presentation templates.")

    # Directory arguments
    parser.add_argument('--img-dnld-dir', type=str, required=True, help="Main image scratch directory.")
    parser.add_argument('--img-eval-dir', type=str, required=True, help="Evaluation image work directory.")
    parser.add_argument('--img-circ-dir', type=str, required=True, help="Circulation image scratch directory.")
    parser.add_argument('--img-sie-dir', type=str, required=True, help="SIE image scratch directory.")
    parser.add_argument('--img-web-dir', type=str, required=True, help="Web image directory.")
    parser.add_argument('--img-diag-dir', type=str, required=True, help="Nino3.4 plot directory.")
    parser.add_argument('--img-wami-dir', type=str, required=True, help="WAMI plot directory.")
    parser.add_argument('--output-file', '-o', type=str, default=None, help="Path to save output briefing.")

    # String parameters
    parser.add_argument('positional_args', nargs='*', help="Parameters for text replacement and filename generation.")

    parsed, _ = parser.parse_known_args()
    args = parsed.positional_args

    now = datetime.now()
    month = now.month

    # Image directory variables
    img_desc_vars = {
        'IMG_DNLD_DIR': parsed.img_dnld_dir,
        'IMG_EVAL_DIR': parsed.img_eval_dir,
        'IMG_CIRC_DIR': parsed.img_circ_dir,
        'IMG_SIE_DIR': parsed.img_sie_dir,
        'IMG_WEB_DIR': parsed.img_web_dir,
        'IMG_DIAG_DIR': parsed.img_diag_dir,
        'IMG_WAMI_DIR': parsed.img_wami_dir,
    }

    # Bash input parameters
    if len(args) >= 19:
        img_desc_vars.update({
            'yyyy': args[0],
            'mm': args[1],
            'mmstring': args[2],
            'yyyym1': args[3],
            'yyyym2': args[4],
            'mm1': args[5],
            'mm2': args[6],
            'mm1str': args[7],
            'yyyym2': args[8],
            'mm2str': args[9],
            'sm1forey': args[10],
            'sm1forem': args[11],
            'sm1foremstr': args[12],
            'seam1': args[13],
            'sea': args[14],
            'dayofbriefing': args[15],
            'nddm1': args[16],
            'mm6str': args[17],
            'yyyym6': args[18],
        })

        # Text substitutions
        replacements = {
            'MM1STR MM1Y': f"{img_desc_vars['mm1str']} {img_desc_vars['yyyym1']}",
            'MM6STR MM6Y': f"{img_desc_vars['mm6str']} {img_desc_vars['yyyym6']}",
            'MM1YMM1NDDM1': f"{img_desc_vars['yyyym1']}{img_desc_vars['mm1']}{img_desc_vars['nddm1']}",
            'MM2STR MM2Y': f"{img_desc_vars['mm2str']} {img_desc_vars['yyyym2']}",
            'SeAM1': f"{img_desc_vars['seam1']}",
            'SM1ForeM SM1ForeY': f"{img_desc_vars['sm1foremstr']} {img_desc_vars['sm1forey']}",
            'ThisM ThisY': f"{img_desc_vars['mmstring']} {img_desc_vars['yyyy']}",
            'SEA': f"{img_desc_vars['sea']}",
            'GLOBAL: t2m': 'GLOBAL: na schifezza',
            'EUROPE: t2m': 'EUROPE: nu purpo',
            'WET': 'nebbia in val padana',
            'DRY': 'martini',
            'HOT': 'afa',
            'COLD': 'magari',
            'today': f"{img_desc_vars['dayofbriefing']}/{img_desc_vars['mm']}/{img_desc_vars['yyyy']}",
        }

        # Load template based on season
        script_dir = Path(__file__).resolve().parent
        template_summer = script_dir.parent / 'pptx' / 'template_summer.pptx'
        template_winter = script_dir.parent / 'pptx' / 'template_winter.pptx'
        template_path = template_summer if (4 <= month < 10) else template_winter
        template_name = str(template_path)
        prs = Presentation(template_name)

        # Replace text
        replace_text_in_presentation(prs, replacements)

        # Replace images using named shape lookup
        image_mappings = get_named_slide_image_mapping(month)
        for slide_idx, shape_map in image_mappings.items():
            if slide_idx < len(prs.slides):
                replace_named_images_on_slide(prs.slides[slide_idx], shape_map, img_desc_vars)

        # Output presentation pptx
        if parsed.output_file:
            output_path = Path(parsed.output_file).resolve()
        else:
            month_name = now.strftime('%B')
            year_name = now.strftime('%Y')
            output_path = Path(parsed.img_dnld_dir) / f"{month_name}_{year_name}.pptx"

        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))
        print(f"Presentation successfully created at: {output_path}")
